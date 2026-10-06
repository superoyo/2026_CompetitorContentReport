# -*- coding: utf-8 -*-
"""The Analysis Report's slides, drawn into build_slides.py's deck.

build_slides.py calls render(globals()) when the report carries `report`
(see analyse_period.py): four parts, each opening with a divider, every
section a table of counted figures beside the model's reading of them. A
section the data cannot support says what is missing instead of showing
an empty table.
"""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PARTS = [
    ("1", "Content Performance Analysis", "การวิเคราะห์ประสิทธิภาพเนื้อหา",
     "เนื้อหาแบบไหนที่โดนใจผู้ติดตามมากที่สุด — รูปแบบโพสต์ หมวดหมู่เนื้อหา และองค์ประกอบของข้อความ"),
    ("2", "Temporal & Scheduling Analysis", "การวิเคราะห์ด้านเวลาและความถี่",
     "ช่วงเวลาทองและจังหวะที่เหมาะสมในการโพสต์ — วัน เวลา ความถี่ และอายุของโพสต์"),
    ("3", "Audience Behavior & Quality Analysis", "การวิเคราะห์พฤติกรรมและคุณภาพการมีส่วนร่วม",
     "สัดส่วนและนัยของตัวเลข Engagement — คุณภาพ อารมณ์ของผู้ชม และ Engagement Rate"),
    ("4", "Correlation & Trend Analysis", "การวิเคราะห์ความสัมพันธ์และแนวโน้ม",
     "ความสัมพันธ์ของตัวแปร ทิศทางการเติบโต และโพสต์ที่ผิดปกติ"),
]
REACT_TH = {"like": "Like", "love": "Love", "care": "Care", "haha": "Haha",
            "wow": "Wow", "sad": "Sad", "angry": "Angry"}


def render(g):
    R = g["P"]["report"]
    S, SEC = R["stats"], R["sections"]
    hx, add_slide, set_bg, add_text, add_rect = g["hx"], g["add_slide"], g["set_bg"], g["add_text"], g["add_rect"]
    add_ring, add_badge = g["add_ring"], g["add_badge"]
    WHITE, DARK, MUTED, ACCENT, TEXT_DARK, RING = (g["WHITE"], g["DARK"], g["MUTED"], g["ACCENT"],
                                                   g["TEXT_DARK"], g["RING"])
    HEAD, BODY = g["HEAD_FONT"], g["BODY_FONT"]
    NAME, COLOR, SPAN = g["NAME"], g["COLOR"], g["SPAN"]
    brands = [b["key"] for b in S["brands"]]
    short = {k: NAME.get(k, k).replace(" Thailand", "") for k in brands}
    avail = S.get("available") or {}

    def num(v, nd=0):
        if v is None:
            return "—"
        if isinstance(v, float) and nd:
            return f"{v:,.{nd}f}"
        return f"{int(round(v)):,}"

    def table(s, x, y, w, header, rows, widths=None, best_cols=None, font=10, row_h=0.34):
        """rows: [[cell...]]; best_cols: highlight the row maximum among these
        column indexes (cells carry (text, value) tuples there)."""
        nr, nc = len(rows) + 1, len(header)
        shape = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
        t = shape.table
        widths = widths or [w / nc] * nc
        for i, cw in enumerate(widths):
            t.columns[i].width = Inches(cw)
        for r in range(nr):
            t.rows[r].height = Inches(row_h)

        def put(cell, text, bold=False, color=TEXT_DARK, fill=None, align=PP_ALIGN.CENTER, size=font):
            cell.text = ""
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = align
            run = p.add_run()
            run.text = str(text)
            run.font.size = Pt(size); run.font.bold = bold
            run.font.name = BODY; run.font.color.rgb = color
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill or WHITE

        for c, h in enumerate(header):
            put(t.cell(0, c), h, bold=True, color=WHITE, fill=hx("#0B2545"),
                align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER, size=font - 0.5)
        for r, row in enumerate(rows, 1):
            vals = {}
            if best_cols:
                vals = {c: row[c][1] for c in best_cols
                        if isinstance(row[c], tuple) and row[c][1] is not None}
            top = max(vals.values()) if vals else None
            stripe = hx("#F6F8FC") if r % 2 == 0 else WHITE
            for c, cellv in enumerate(row):
                text = cellv[0] if isinstance(cellv, tuple) else cellv
                hit = top is not None and c in vals and vals[c] == top and len(vals) > 1
                put(t.cell(r, c), text, bold=hit or c == 0,
                    color=hx("#9A5B00") if hit else TEXT_DARK,
                    fill=hx("#FFF1D6") if hit else stripe,
                    align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
        return shape

    def head(s, sid, title=None):
        sec = SEC.get(sid) or {}
        add_text(s, 0.6, 0.38, 12.1, 0.55, title or sec.get("title", ""), size=23,
                 color=hx("#0B2545"), bold=True, font=HEAD)
        if sec.get("headline"):
            add_text(s, 0.6, 0.95, 12.1, 0.5, sec["headline"], size=12.5, color=hx("#9A5B00"), bold=True)

    def insights(s, sids, x=8.75, y=1.6, w=3.98, h=5.4):
        add_rect(s, x, y, w, h, fill=hx("#F4F7FB"), radius=0.04)
        add_rect(s, x, y, 0.08, h, fill=ACCENT)
        add_text(s, x + 0.25, y + 0.15, w - 0.4, 0.35, "ข้อค้นพบ", size=12.5, color=hx("#0B2545"),
                 bold=True, font=HEAD)
        lines = []
        for sid in sids:
            for it in (SEC.get(sid) or {}).get("insights") or []:
                lines.append(("• " + it, {"space_before": 5}))
        add_text(s, x + 0.25, y + 0.55, w - 0.45, h - 0.7, lines or [("—", {})], size=10.5,
                 color=TEXT_DARK, ls=1.08)

    def missing(s, text, x=0.6, y=1.6, w=7.9, h=1.3):
        add_rect(s, x, y, w, h, fill=hx("#F1F3F6"), radius=0.05, line_color=hx("#D5DBE3"))
        add_text(s, x + 0.25, y + 0.15, w - 0.5, h - 0.3, [("ข้อมูลไม่พอสำหรับหัวข้อนี้", {"bold": True, "size": 12}),
                                                           (text, {"size": 10.5, "space_before": 4})],
                 color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

    def white(sid=None, title=None):
        s = add_slide(); set_bg(s, WHITE)
        head(s, sid, title)
        return s

    def divider(i):
        no, en, th, desc = PARTS[i]
        s = add_slide(); set_bg(s, DARK)
        add_ring(s, 13.2, 6.8, 2.8, RING, 1.6); add_ring(s, -0.2, -0.3, 1.4, RING, 1.6)
        add_text(s, 1.0, 1.9, 6, 0.6, "PART " + no, size=20, color=ACCENT, bold=True, font=HEAD)
        add_text(s, 1.0, 2.5, 11.3, 0.9, en, size=38, color=WHITE, bold=True, font=HEAD)
        add_text(s, 1.0, 3.45, 11.3, 0.6, th, size=22, color=hx("#CFE7F0"), font=HEAD)
        add_rect(s, 1.0, 4.25, 0.9, 0.07, fill=ACCENT)
        add_text(s, 1.0, 4.5, 10.5, 0.8, desc, size=14, color=hx("#9FC9DE"))

    n_brands_h = min(0.34, 4.9 / max(len(brands) + 2, 1))

    # ---------------- PART 1
    divider(0)
    ct = S["content_type"]
    types = list(ct["labels"])
    s = white("content_type")
    rows = []
    for k in brands + ["_group"]:
        src = ct["group"] if k == "_group" else ct["brands"].get(k) or {}
        row = ["รวมทั้งกลุ่ม" if k == "_group" else short[k]]
        for t in types:
            c = src.get(t)
            row.append(("%s (%d)" % (num(c["avg"]), c["n"]), c["avg"]) if c else ("—", None))
        rows.append(row)
    table(s, 0.6, 1.6, 7.9, ["เพจ · Engagement เฉลี่ย/โพสต์ (จำนวนโพสต์)"] + [ct["labels"][t] for t in types],
          rows, widths=[2.3] + [5.6 / max(len(types), 1)] * len(types),
          best_cols=list(range(1, len(types) + 1)), row_h=n_brands_h)
    if not avail.get("album_link"):
        add_text(s, 0.6, 6.95, 7.9, 0.3, "อัลบั้มรูปและลิงก์แยกได้เฉพาะเดือนที่ดึงข้อมูลหลังการอัปเดตระบบ "
                 "เดือนเก่านับรวมอยู่ในรูปภาพ/อื่น ๆ", size=9, color=MUTED, italic=True)
    insights(s, ["content_type"])

    s = white("pillar")
    if S.get("pillar"):
        pl = S["pillar"]
        total_n = sum(v["n"] for v in pl["group"].values()) or 1
        rows = [[pl["labels"][pid], "%d (%.0f%%)" % (v["n"], 100.0 * v["n"] / total_n),
                 (num(v["avg"]), v["avg"]), (num(v["avg_comments"], 1), v["avg_comments"]),
                 (num(v["avg_shares"], 1), v["avg_shares"])] for pid, v in pl["group"].items()]
        table(s, 0.6, 1.6, 7.9, ["หมวดหมู่เนื้อหา", "โพสต์", "Engagement เฉลี่ย", "คอมเมนต์เฉลี่ย", "แชร์เฉลี่ย"],
              rows, widths=[2.7, 1.3, 1.4, 1.25, 1.25], row_h=0.4)
        best = []
        for k in brands:
            cand = [(pid, v) for pid, v in (pl["brands"].get(k) or {}).items() if v["n"] >= 2]
            if cand:
                pid, v = max(cand, key=lambda x: x[1]["avg"])
                best.append("%s: %s" % (short[k], pl["labels"][pid]))
        if best:
            add_text(s, 0.6, 5.35, 7.9, 1.6, [("หมวดที่ได้ Engagement เฉลี่ยสูงสุดของแต่ละเพจ (อย่างน้อย 2 โพสต์)",
                                               {"bold": True, "size": 10.5})]
                     + [("• " + b, {"size": 9.5}) for b in best], color=TEXT_DARK)
        add_text(s, 0.6, 7.0, 7.9, 0.3, "จัดหมวดหมู่จากแคปชันโดย AI", size=9, color=MUTED, italic=True)
    else:
        missing(s, "จัดหมวดหมู่เนื้อหาไม่สำเร็จในรอบนี้ ลองสร้างรายงานใหม่หรือเลือกโมเดลอื่น")
    insights(s, ["pillar"])

    s = white("copywriting")
    rows = []
    for factor, buckets in S["copywriting"].items():
        first = True
        best = max((v["avg"] for v in buckets.values()), default=None)
        for b, v in buckets.items():
            hit = v["avg"] == best and len(buckets) > 1
            rows.append([factor if first else "", b, str(v["n"]),
                         ("★ " if hit else "") + num(v["avg"]), num(v["avg_comments"], 1)])
            first = False
    table(s, 0.6, 1.6, 7.9, ["ปัจจัย", "กลุ่ม", "โพสต์", "Engagement เฉลี่ย", "คอมเมนต์เฉลี่ย"], rows,
          widths=[1.7, 2.3, 1.0, 1.6, 1.3], row_h=min(0.36, 5.2 / (len(rows) + 1)))
    add_text(s, 0.6, 7.0, 7.9, 0.3, "ไม่มียอดคลิกจากข้อมูลสาธารณะ จึงวัดผลของ CTA ได้จาก Engagement และคอมเมนต์เท่านั้น",
             size=9, color=MUTED, italic=True)
    insights(s, ["copywriting"])

    # ---------------- PART 2
    divider(1)
    s = white("timing")
    days = ["จันทร์", "อังคาร", "พุธ", "พฤหัส", "ศุกร์", "เสาร์", "อาทิตย์"]
    rows = []
    for k in brands + ["_group"]:
        src = S["day"]["group"] if k == "_group" else S["day"]["brands"].get(k) or {}
        rows.append(["รวมทั้งกลุ่ม" if k == "_group" else short[k]]
                    + [(num(src[d]["avg"]), src[d]["avg"]) if d in src else ("—", None) for d in days])
    h_tab = n_brands_h * (len(rows) + 1)
    table(s, 0.6, 1.6, 7.9, ["Engagement เฉลี่ย/โพสต์"] + days, rows, widths=[2.0] + [5.9 / 7] * 7,
          best_cols=list(range(1, 8)), row_h=n_brands_h, font=9.5)
    y = 1.6 + h_tab + 0.25
    if S.get("hour"):
        hg = S["hour"]["group"]
        slots = list(hg)
        table(s, 0.6, y, 7.9, ["ช่วงเวลา (ทั้งกลุ่ม)"] + slots,
              [["Engagement เฉลี่ย (โพสต์)"] + [("%s (%d)" % (num(hg[x]["avg"]), hg[x]["n"]), hg[x]["avg"])
                                              for x in slots]],
              widths=[2.3] + [5.6 / len(slots)] * len(slots), best_cols=list(range(1, len(slots) + 1)),
              row_h=0.34, font=9.5)
    else:
        missing(s, "เวลาโพสต์ (ชั่วโมง) มีเฉพาะเดือนที่ดึงหลังการอัปเดตระบบ ดึงข้อมูลเดือนเหล่านี้ใหม่เพื่อดูช่วงเวลาทอง",
                y=min(y, 5.9), h=1.0)
    insights(s, ["timing"])

    s = white("frequency", "2.2–2.3 ความถี่ในการโพสต์ และอายุของโพสต์")
    fr = S["frequency"]
    rows = [[short[k], str(v["posts"]), num(v["per_week"], 1), num(v["avg_per_post"]),
             ("%.2f" % v["r_posts_vs_avg"]) if v["r_posts_vs_avg"] is not None else "—"]
            for k, v in fr["brands"].items()]
    table(s, 0.6, 1.6, 7.9, ["เพจ", "โพสต์", "โพสต์/สัปดาห์", "Engagement/โพสต์", "r (ความถี่ กับ Eng/โพสต์)"],
          rows, widths=[2.2, 1.0, 1.4, 1.6, 1.7], row_h=n_brands_h)
    add_text(s, 0.6, 1.6 + n_brands_h * (len(rows) + 1) + 0.1, 7.9, 0.4,
             "r รวมทุกเพจ = %s (%d สัปดาห์) · r คำนวณเมื่อเพจมีโพสต์อย่างน้อย 4 สัปดาห์" % (
                 ("%.2f" % fr["pooled_r"]) if fr["pooled_r"] is not None else "—", fr["pooled_weeks"]),
             size=9.5, color=MUTED)
    missing(s, "อายุของโพสต์ต้องเก็บยอดของโพสต์เดียวกันหลายครั้งตามเวลา (เช่น ทุก 6 ชั่วโมงในสัปดาห์แรก) "
               "ข้อมูลปัจจุบันเก็บครั้งเดียวต่อเดือน", y=5.75, h=1.2)
    insights(s, ["frequency", "lifespan"])

    # ---------------- PART 3
    divider(2)
    s = white("quality", "3.1 + 3.3 คุณภาพของ Engagement และ Engagement Rate")
    q = S["quality"]
    rows = [[short[k], num(v["likes_per_post"]), (num(v["comments_per_post"], 1), v["comments_per_post"]),
             (num(v["shares_per_post"], 1), v["shares_per_post"]), "%.1f%%" % v["comment_pct"],
             "%.1f%%" % v["share_pct"], num(v["followers"]),
             ("%.3f%%" % v["er_per_post"]) if v["er_per_post"] is not None else "—"] for k, v in q.items()]
    table(s, 0.6, 1.6, 7.9, ["เพจ", "ไลก์/โพสต์", "คอมเมนต์/โพสต์", "แชร์/โพสต์", "%คอมเมนต์", "%แชร์",
                             "ผู้ติดตาม", "ER/โพสต์"], rows,
          widths=[1.6, 0.95, 1.05, 0.95, 0.85, 0.75, 0.95, 0.8], row_h=n_brands_h, font=9.5)
    add_text(s, 0.6, 6.7, 7.9, 0.5, "ER/โพสต์ = Engagement ของโพสต์ ÷ ผู้ติดตามของเพจ ณ เดือนนั้น × 100 "
             "(ไม่มี Reach จากข้อมูลสาธารณะ) · ไม่มียอดคลิกลิงก์", size=9, color=MUTED, italic=True)
    insights(s, ["quality", "er"])

    s = white("sentiment")
    if S.get("sentiment"):
        keys = [r for r in REACT_TH if any(r in v for v in S["sentiment"].values())]
        rows = [[short[k]] + [("%.1f%%" % v[r]) if r in v else "—" for r in keys] + [num(v["_total"])]
                for k, v in S["sentiment"].items()]
        table(s, 0.6, 1.6, 7.9, ["เพจ"] + [REACT_TH[r] for r in keys] + ["Reaction รวม"], rows,
              widths=[1.9] + [5.0 / len(keys)] * len(keys) + [1.0], row_h=n_brands_h, font=9.5)
    else:
        missing(s, "สัดส่วน Reaction (Like, Love, Haha, Wow, Sad, Angry) มีเฉพาะเดือนที่ดึงหลังการอัปเดตระบบ "
                   "ดึงข้อมูลเดือนเหล่านี้ใหม่เพื่อดูอารมณ์ของผู้ชม")
    insights(s, ["sentiment"])

    # ---------------- PART 4
    divider(3)
    s = white("correlation")
    rows = [[c["pair"], ("%.2f" % c["r"]) if c["r"] is not None else "—", str(c["n"]), c["strength"]]
            for c in S["correlation"]]
    table(s, 0.6, 1.6, 7.9, ["คู่ตัวแปร", "r (Spearman)", "n", "ระดับความสัมพันธ์"], rows,
          widths=[3.6, 1.1, 0.8, 2.4], row_h=0.42)
    add_text(s, 0.6, 1.6 + 0.42 * (len(rows) + 1) + 0.15, 7.9, 0.8,
             "r ใกล้ +1 = ไปทางเดียวกัน · ใกล้ −1 = สวนทาง · ใกล้ 0 = ไม่สัมพันธ์ "
             "ความสัมพันธ์ไม่ได้แปลว่าเป็นเหตุเป็นผล · ไม่มี Reach/คลิกจึงทดสอบคู่ที่ใช้ตัวแปรเหล่านี้ไม่ได้",
             size=9.5, color=MUTED)
    insights(s, ["correlation"])

    s = white("growth")
    gm = S["growth"]["monthly"]
    months = [r["month"] for r in next(iter(gm.values()), [])]
    import month_util
    mlab = ["%s %d" % (month_util.TH_ABBR[int(m[5:7]) - 1], (int(m[:4]) + 543) % 100) for m in months]
    rows = []
    for k, rs in gm.items():
        row = [short[k]]
        for r in rs:
            row.append(num(r["total"]) + ("" if r["mom"] is None else "\n%s%.0f%%" % ("+" if r["mom"] >= 0 else "", r["mom"])))
        rows.append(row)
    table(s, 0.6, 1.6, 7.9, ["Engagement รวม (MoM)"] + mlab, rows,
          widths=[2.2] + [5.7 / max(len(months), 1)] * len(months), row_h=max(n_brands_h, 0.42), font=9.5)
    wk = S["growth"]["weekly"]
    if wk:
        up = sum(1 for w in wk if (w["wow"] or 0) > 0)
        add_text(s, 0.6, 6.65, 7.9, 0.5, "รายสัปดาห์ (ทั้งกลุ่ม): %d สัปดาห์ · สัปดาห์ที่ Engagement รวมสูงขึ้นจากสัปดาห์ก่อน %d สัปดาห์"
                 % (len(wk), up), size=9.5, color=MUTED)
    insights(s, ["growth"])

    s = white("anomalies")
    an = S.get("anomalies") or []
    reasons = R.get("anomaly_reasons") or {}
    if an:
        rows = [[short.get(a["brand"], a["name"]) + "\n" + a["date"],
                 "%s\n%s" % (a["direction"], a["type"]), num(a["total"]),
                 "×%s" % a["x_median"],
                 (reasons.get(a["id"]) or g["truncate"](a["text"], 90))] for a in an[:8]]
        table(s, 0.6, 1.6, 12.1, ["เพจ · วันที่", "ประเภท", "Engagement", "เทียบค่ากลาง", "สาเหตุที่น่าจะเป็น (AI)"],
              rows, widths=[1.9, 1.5, 1.2, 1.1, 6.4], row_h=0.6, font=9)
        add_text(s, 0.6, 7.0, 12.1, 0.3, "ผิดปกติ = log(Engagement) ห่างจากค่าเฉลี่ยของเพจตัวเองเกิน 2 SD · "
                 "สาเหตุเป็นข้อสันนิษฐาน (การยิงแอดยืนยันไม่ได้จากข้อมูลสาธารณะ ยกเว้นป้าย Paid partnership)",
                 size=9, color=MUTED, italic=True)
    else:
        missing(s, "ไม่พบโพสต์ที่ผิดปกติเกินเกณฑ์ (2 SD) หรือมีโพสต์ต่อเพจน้อยเกินไป")
        insights(s, ["anomalies"])

    # ---------------- recommendations
    s = add_slide(); set_bg(s, DARK)
    add_ring(s, 13.4, 7.6, 2.6, RING, 1.6)
    add_text(s, 0.8, 0.6, 11.5, 0.6, "สรุปและข้อเสนอแนะ", size=28, color=WHITE, bold=True, font=HEAD)
    add_text(s, 0.8, 1.2, 11.5, 0.4, "จากการวิเคราะห์ทั้ง 4 ด้าน · " + SPAN, size=13, color=hx("#9FC9DE"))
    y = 1.9
    for i, rtext in enumerate(R.get("recommendations") or [], 1):
        add_badge(s, 0.8, y + 0.02, 0.42, ACCENT, str(i), tc=DARK, fs=12)
        add_text(s, 1.4, y, 11.2, 0.7, rtext, size=13, color=hx("#E3EEF4"), ls=1.1)
        y += 0.72
